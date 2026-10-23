"""Test the production Office stack and XML linkage with synthetic documents."""
import io
import json
import os
from pathlib import Path
import re
import subprocess
import tempfile

from docx import Document
from docx.shared import Inches as WordInches
from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as SheetImage
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfReader

VERSION = '2.12.7+dfsg+really2.9.14-2.1+deb13u3+webui1'


def run(*args, **kwargs):
    return subprocess.run(args, check=True, capture_output=True, text=True, timeout=60, **kwargs)


def main():
    assert os.getuid() == 11002, 'Use the real production user'
    version = run('dpkg-query', '-W', '-f=${Version}', 'libxml2').stdout.strip()
    assert version == VERSION
    assert etree.LIBXML_VERSION == etree.LIBXML_COMPILED_VERSION == (2, 9, 14)
    linkage = run('ldd', etree.__file__).stdout
    assert 'libxml2.so.2 =>' in linkage and 'not found' not in linkage
    mapped = sorted({line.split()[-1] for line in Path('/proc/self/maps').read_text().splitlines()
                     if '/libxml2.so.' in line})
    assert len(mapped) == 1, mapped
    assert '/usr/lib/' in mapped[0] or '/lib/' in mapped[0]
    assert '.libs/' not in mapped[0], 'Do not retain a second bundled libxml2'
    office_linkage = []
    for module in Path('/usr/lib/libreoffice/program').glob('*.so'):
        result = subprocess.run(['readelf', '-d', str(module)], capture_output=True, text=True, timeout=10)
        if 'Shared library: [libxml2.so.2]' in result.stdout:
            resolved = run('ldd', str(module)).stdout
            assert 'not found' not in resolved
            office_linkage.append(module.name)
    assert office_linkage, 'No Office consumer of system libxml2 found'

    source_review = json.loads(Path('/usr/local/share/libxml2-backport/source-review.json').read_text())
    bindings = subprocess.run(['dpkg-query', '-W', '-f=${db:Status-Status}', 'python3-libxml2'], capture_output=True, text=True)
    assert bindings.stdout.strip() != 'installed'
    checks = []
    with tempfile.TemporaryDirectory(prefix='webui-office-') as directory:
        root = Path(directory)
        image_path = root / 'image.png'
        Image.new('RGB', (160, 80), (30, 90, 170)).save(image_path)
        doc = Document()
        doc.add_heading('WebUI XML Word', 0)
        doc.add_paragraph('Texte accentué : été, Noël, caractères & symboles.')
        doc.add_table(rows=2, cols=2).cell(0, 0).text = 'Table marker'
        doc.add_picture(str(image_path), width=WordInches(1))
        doc.add_page_break()
        doc.add_heading('Second page marker', 1)
        word = root / 'word.docx'
        doc.save(word)

        workbook = Workbook()
        sheet = workbook.active
        sheet['A1'] = 'WebUI XML Calc'
        sheet['A2'], sheet['A3'], sheet['A4'] = 'First', 'Second', 'Total'
        sheet['B2'], sheet['B3'], sheet['B4'] = 12, 30, '=SUM(B2:B3)'
        sheet.column_dimensions['A'].width = 28
        sheet.add_image(SheetImage(str(image_path)), 'D2')
        sheet.print_area = 'A1:G10'
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = sheet.page_setup.fitToHeight = 1
        excel = root / 'sheet.xlsx'
        workbook.save(excel)

        slides = Presentation()
        for index, text in enumerate(['WebUI XML Impress', 'Second slide marker']):
            slide = slides.slides.add_slide(slides.slide_layouts[5])
            slide.shapes.title.text = text
            box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(5), Inches(1))
            box.text_frame.text = 'Texte accentué et symboles : été & Noël'
            slide.shapes.add_picture(str(image_path), Inches(1), Inches(4), width=Inches(2))
        powerpoint = root / 'slides.pptx'
        slides.save(powerpoint)

        def convert(source, target):
            number = len(checks)
            output = root / f'out-{number}'
            output.mkdir()
            profile = (root / f'profile-{number}').as_uri()
            result = run('libreoffice', f'-env:UserInstallation={profile}', '--headless',
                         '--nologo', '--nodefault', '--nofirststartwizard', '--convert-to', target,
                         '--outdir', str(output), str(source))
            path = output / (source.stem + '.' + target.split(':')[0])
            assert path.is_file() and path.stat().st_size > 100, (result.stdout, result.stderr)
            checks.append({'source': source.suffix, 'target': target, 'bytes': path.stat().st_size})
            return path

        def pdf(source, markers, pages, images=False):
            path = convert(source, 'pdf')
            reader = PdfReader(path)
            assert len(reader.pages) == pages, (source, len(reader.pages))
            text = re.sub(r'\s+', ' ', ' '.join(page.extract_text() or '' for page in reader.pages))
            for marker in markers:
                assert marker in text, (source, marker, text)
            if images:
                assert sum(len(page.images) for page in reader.pages) >= 1

        pdf(word, ['WebUI XML Word', 'Second page marker'], 2, images=True)
        odt = convert(word, 'odt')
        pdf(odt, ['WebUI XML Word', 'Second page marker'], 2, images=True)
        word_roundtrip = convert(odt, 'docx')
        reopened = Document(word_roundtrip)
        assert any('été' in p.text for p in reopened.paragraphs)
        assert reopened.tables and len(reopened.inline_shapes) == 1
        assert 'Table marker' in reopened.tables[0].cell(0, 0).text
        assert reopened.paragraphs[0].style.name != 'Normal'
        legacy_word = convert(word, 'doc:MS Word 97')
        pdf(legacy_word, ['WebUI XML Word', 'Second page marker'], 2, images=True)

        pdf(excel, ['WebUI XML Calc', '42'], 1, images=True)
        ods = convert(excel, 'ods')
        pdf(ods, ['WebUI XML Calc', '42'], 1, images=True)
        excel_roundtrip = convert(ods, 'xlsx')
        assert load_workbook(excel_roundtrip, data_only=True).active['B4'].value == 42
        assert load_workbook(excel_roundtrip, data_only=False).active['B4'].value.startswith('=')
        legacy_excel = convert(excel, 'xls:MS Excel 97')
        pdf(legacy_excel, ['WebUI XML Calc', '42'], 1, images=True)

        pdf(powerpoint, ['WebUI XML Impress', 'Second slide marker'], 2, images=True)
        odp = convert(powerpoint, 'odp')
        pdf(odp, ['WebUI XML Impress', 'Second slide marker'], 2, images=True)
        slides_roundtrip = Presentation(convert(odp, 'pptx'))
        assert len(slides_roundtrip.slides) == 2
        assert any('été' in shape.text for slide in slides_roundtrip.slides
                   for shape in slide.shapes if shape.has_text_frame)
        legacy_slides = convert(powerpoint, 'ppt:MS PowerPoint 97')
        pdf(legacy_slides, ['WebUI XML Impress', 'Second slide marker'], 2, images=True)
        assert len(checks) == 18

        # LibreOffice may deliberately import unknown XML as plain text. Test
        # strict XML rejection through the patched lxml parser, not that fallback.
        strict_rejected = False
        try:
            etree.fromstring(b'<root><unclosed>', etree.XMLParser(
                recover=False, resolve_entities=False, no_network=True))
        except etree.XMLSyntaxError:
            strict_rejected = True
        assert strict_rejected, 'Strict XML parsing must reject malformed input'

    print(json.dumps({'package_version': version, 'production_uid': os.getuid(),
                      'lxml_version': '.'.join(map(str, etree.LXML_VERSION)),
                      'lxml_runtime_libxml2': list(etree.LIBXML_VERSION),
                      'lxml_system_linkage': mapped, 'office_xml_consumers': office_linkage,
                      'source_review': source_review, 'native_python_bindings_installed': False,
                      'conversions': checks, 'pdf_text_pages_images': 'passed',
                      'word_styles_tables_unicode_roundtrip': 'passed', 'calc_formula_cached_value': 42,
                      'slides_unicode_roundtrip': 'passed', 'strict_xml_rejected_in_lxml': strict_rejected,
                      'network': 'disabled by Docker wrapper'}, indent=2))


if __name__ == '__main__':
    main()
